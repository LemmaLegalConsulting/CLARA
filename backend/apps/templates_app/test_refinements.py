"""A reviewer's refinement file turns a Word original into a finished filing.

The Cleveland originals mix fill-ins with editorial notes -- "[written/oral]",
"[X] days", "[OR]" between two certificates, a "Case Caption" marker -- and a
prepared template printed them in the filing. These tests pin down the
mechanism that lets a reviewer say what each passage means, and the export
behaviour that lets the maintained, formatted wording reach the document.
"""

import io
import re
import tempfile
from pathlib import Path
from types import SimpleNamespace

import yaml
from django.test import SimpleTestCase, TestCase
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from apps.ai.services import ConstrainedDraftingService, GenerationContext
from apps.core.content_library import organization_content_library_dir
from apps.exporting.services import _block_revision, _render_docx_template
from apps.templates_app.field_questions import KIND_NARRATIVE, field_applies, template_field_requests
from apps.templates_app.ingestion import discover_blocks, ingest_docx
from apps.templates_app.jinja_filters import answered_items, currency
from apps.templates_app.refinements import RefinementError, load_refinements
from apps.templates_app.template_variables import (
    field_defaults,
    normalize_docxtpl_blocks,
    signature_lines,
    template_choice_values,
    template_field_values,
)


def highlighted(paragraph, text):
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def write_original(path: Path):
    document = Document()
    marker = document.add_paragraph()
    highlighted(marker, "Case Caption - Defendant’s Motion")
    motion = document.add_paragraph("Now comes Defendant [DEFENDANT]. The lease is ")
    motion.add_run("[written/oral]").bold = True
    motion.add_run(". Discovery needs [X] days, [etc.].")
    document.add_heading("CERTIFICATE OF SERVICE", level=1)
    document.add_paragraph("I served it on [DATE] by email to [PLAINTIFF’S EMAIL].")
    document.add_paragraph("[OR]")
    document.add_paragraph("I served it on [DATE] by United States mail to [PLAINTIFF’S ADDRESS].")
    document.add_heading("STATEMENT OF THE CASE", level=1)
    document.add_paragraph("Defendant and [other occupants] live at the premises.")
    document.save(path)


REFINEMENTS = {
    "schema_version": 1,
    "slug": "sample-motion",
    "title": "Defendant's Sample Motion",
    "caption": {
        "marker": "Case Caption - Defendant’s Motion",
        "court": ["IN THE CLEVELAND MUNICIPAL COURT"],
        "left": ["{{ plaintiff }},", "Plaintiff,", "v.", "{{ defendant }},", "Defendant."],
        "right": ["CASE NO. {{ case_number }}"],
        "title": "DEFENDANT’S MOTION",
    },
    "blocks": {"document-body": {"label": "Motion", "ai_latitude": "locked"}},
    "choices": [
        {"name": "lease_form", "label": "Lease", "options": [{"value": "written", "label": "Written"}, {"value": "oral", "label": "Oral"}], "default": "written"},
    ],
    "fields": {
        "discovery_days": {"label": "Days needed for discovery", "question": "How many days does discovery need?"},
        "other_occupants": {"label": "Other occupants", "type": "list", "kind": "value"},
        "plaintiff_email": {"label": "Email for service", "when": {"service_method": "email"}},
    },
    "edits": [
        {"find": "The lease is [written/oral].", "replace": 'The lease is {% if lease_form == "oral" %}oral{% else %}written{% endif %}.'},
        {"find": "needs [X] days, [etc.].", "replace": "needs {{ fields.discovery_days }} days."},
        {"find": "Defendant and [other occupants] live", "replace": '{{ (["Defendant"] + fields.other_occupants|answered_items)|comma_and_list }} live'},
    ],
}


class RefinedIngestionTests(SimpleTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.source = root / "original_templates" / "Sample Motion.docx"
        self.source.parent.mkdir()
        write_original(self.source)
        self.refinements = root / "template-refinements" / "sample-motion.yaml"
        self.refinements.parent.mkdir()
        self.write_refinements(REFINEMENTS)
        self.prepared = root / "document-templates"
        self.snippets = root / "docx-snippets"

    def tearDown(self):
        self.tmp.cleanup()

    def write_refinements(self, data):
        self.refinements.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False))

    def ingest(self, **kwargs):
        manifest_path = ingest_docx(self.source, self.prepared, self.snippets, **kwargs)
        return yaml.safe_load(manifest_path.read_text()), manifest_path.parent / "template.docx"

    def render(self, template_path, **values):
        context = {
            "fields": template_field_values(values.pop("fields", {})),
            "blocks": {},
            "plaintiff": "Pat Landlord",
            "defendant": "Jordan Client",
            "case_number": "2026-CVG-1",
            "lease_form": "written",
            "service_method": "email",
            **values,
        }
        output = Path(self.tmp.name) / "rendered.docx"
        _render_docx_template(str(template_path), context, str(output))
        document = Document(output)
        body = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(body), "\n".join(cells)

    def test_the_editorial_notes_never_reach_the_filing(self):
        manifest, template = self.ingest()
        body, caption = self.render(template, fields={"discovery_days": "45", "other_occupants": ["Sam Client"]}, lease_form="oral")

        self.assertIn("The lease is oral.", body)
        self.assertIn("Discovery needs 45 days.", body)
        self.assertIn("Defendant and Sam Client live at the premises.", body)
        for note in ("[written/oral]", "[X]", "[etc.]", "[OR]", "Case Caption"):
            self.assertNotIn(note, body)
        # Only the chosen certificate prints.
        self.assertIn("by email to", body)
        self.assertNotIn("United States mail", body)
        self.assertIn("Pat Landlord", caption)
        self.assertIn("DEFENDANT’S MOTION", caption)

    def test_an_empty_list_drops_out_of_the_sentence(self):
        _manifest, template = self.ingest()
        body, _caption = self.render(template, fields={"discovery_days": "45"})
        self.assertIn("Defendant live at the premises.", body)

    def test_the_manifest_carries_the_reviewers_decisions(self):
        manifest, _template = self.ingest()

        self.assertEqual(manifest["title"], "Defendant's Sample Motion")
        self.assertIn("fields.discovery_days", manifest["fields"])
        self.assertIn("lease_form", [choice["name"] for choice in manifest["choices"]])
        self.assertIn("service_method", [choice["name"] for choice in manifest["choices"]])
        self.assertEqual(manifest["field_schema"]["discovery_days"]["question"], "How many days does discovery need?")
        self.assertEqual(manifest["source"]["refinements"]["path"], "template-refinements/sample-motion.yaml")
        motion = next(block for block in manifest["blocks"] if block["key"] == "motion")
        self.assertEqual(motion["ai_latitude"], "locked")
        # The reviewer's text preview applies the choice instead of printing "[OR]".
        certificate = next(block for block in manifest["blocks"] if block["key"] == "certificate-of-service")
        self.assertNotIn("[OR]", certificate["body"])
        self.assertIn('{% if service_method == "email" or not service_method %}', certificate["body"])

    def test_a_changed_refinement_file_regenerates_the_package(self):
        manifest, _template = self.ingest()
        changed = {**REFINEMENTS, "title": "Renamed Motion"}
        self.write_refinements(changed)

        manifest, _template = self.ingest()

        self.assertEqual(manifest["title"], "Renamed Motion")

    def test_an_edit_that_no_longer_matches_fails_instead_of_being_skipped(self):
        stale = {**REFINEMENTS, "edits": [{"find": "The lease is [written or oral].", "replace": "x"}]}
        self.write_refinements(stale)

        with self.assertRaisesRegex(RefinementError, "expected 1 match"):
            self.ingest(force=True)

    def test_a_condition_on_an_unknown_choice_fails(self):
        typo = {**REFINEMENTS, "fields": {"plaintiff_email": {"when": {"servce_method": "email"}}}}
        self.write_refinements(typo)

        with self.assertRaisesRegex(RefinementError, "unknown choice"):
            self.ingest(force=True)

    def test_a_choice_without_a_default_takes_its_first_option(self):
        no_default = {**REFINEMENTS, "choices": [{"name": "lease_form", "options": ["written", "oral"]}]}
        self.write_refinements(no_default)

        manifest, _template = self.ingest(force=True)

        lease = next(choice for choice in manifest["choices"] if choice["name"] == "lease_form")
        self.assertEqual(lease["default"], "written")

    def test_two_blocks_relabelled_to_one_key_fail_before_anything_is_written(self):
        colliding = {**REFINEMENTS, "blocks": {"document-body": {"label": "Certificate of Service"}}}
        self.write_refinements(colliding)

        with self.assertRaisesRegex(RefinementError, "more than one block"):
            self.ingest(force=True)

    def test_the_slug_must_name_the_template(self):
        self.write_refinements({**REFINEMENTS, "slug": "another-template"})

        with self.assertRaisesRegex(RefinementError, "does not name"):
            load_refinements(Path(self.tmp.name), "sample-motion")


class SharedCourtSettingsTests(SimpleTestCase):
    """A court's caption and its sitting judge are set once for every template."""

    def test_a_template_extends_the_shared_court_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "template-refinements"
            folder.mkdir()
            (folder / "_court.yaml").write_text(yaml.safe_dump({
                "schema_version": 1,
                "caption": {"court": ["IN THE CLEVELAND MUNICIPAL COURT"], "right": ["{% if fields.judge|answered %}JUDGE {{ fields.judge|upper }}{% endif %}"]},
                "fields": {"judge": {"label": "Judge", "default": "W. Moná Scott"}},
            }, allow_unicode=True))
            (folder / "sample.yaml").write_text(yaml.safe_dump({
                "schema_version": 1, "slug": "sample", "extends": "_court.yaml",
                "caption": {"title": "MOTION"},
                "fields": {"judge": {"question": "Which judge?"}, "magistrate": {"label": "Magistrate"}},
            }))
            first = load_refinements(Path(tmp), "sample")
            self.assertEqual(first.caption["court"], ["IN THE CLEVELAND MUNICIPAL COURT"])
            self.assertEqual(first.caption["title"], "MOTION")
            self.assertEqual(first.field_schema["judge"], {"label": "Judge", "default": "W. Moná Scott", "question": "Which judge?"})

            # A new judge in the shared file changes every template's checksum,
            # so each package regenerates.
            (folder / "_court.yaml").write_text(yaml.safe_dump({"schema_version": 1, "fields": {"judge": {"default": "Someone Else"}}}))
            self.assertNotEqual(load_refinements(Path(tmp), "sample").sha256, first.sha256)

    def test_the_default_fills_the_field_until_someone_answers(self):
        template = SimpleNamespace(metadata={"fieldSchema": {"judge": {"default": "W. Moná Scott"}, "magistrate": {}}})
        self.assertEqual(field_defaults(template), {"judge": "W. Moná Scott"})


class BindingsDoNotLookLikeInstructionsTests(SimpleTestCase):
    def test_comma_and_list_in_a_binding_is_not_a_request_for_a_list(self):
        document = Document()
        document.add_heading("Burden on Parties", level=1)
        document.add_paragraph('{{ (["Defendant"] + fields.occupants|answered_items)|comma_and_list }} would lose the document section.')
        block = discover_blocks(document)[0]
        self.assertFalse(block.expects_list)


class RenderingHelpersTests(SimpleTestCase):
    def test_currency_reads_an_amount_however_it_was_typed(self):
        self.assertEqual(currency("1250"), "$1,250.00")
        self.assertEqual(currency("$1,250"), "$1,250.00")
        self.assertEqual(currency("the entirety of the contract rent"), "the entirety of the contract rent")
        self.assertEqual(currency("[Monthly Rent]"), "[Monthly Rent]")

    def test_answered_items_drops_an_unanswered_list(self):
        self.assertEqual(answered_items("[Other Occupants]"), [])
        self.assertEqual(answered_items(["Sam", ""]), ["Sam"])

    def test_typographic_quotes_survive_outside_tags(self):
        normalized = normalize_docxtpl_blocks("“Premises” and Defendant’s {{ fields[“x”] }}")
        self.assertEqual(normalized, '“Premises” and Defendant’s {{ fields["x"] }}')

    def test_the_signature_names_the_attorney(self):
        block, signature = signature_lines("Avery Advocate", "0012345", "Legal Aid\n216-555-0100")
        self.assertEqual(signature, "/s/ Avery Advocate\nAvery Advocate (0012345)")
        self.assertEqual(block, "/s/ Avery Advocate\nAvery Advocate (0012345)\nLegal Aid\n216-555-0100")

    def test_highlighted_wording_is_kept_unless_turned_off(self):
        template = SimpleNamespace(metadata={"flags": ["include_bond_sentence"], "choices": []})
        self.assertEqual(template_choice_values(template, {}), {"include_bond_sentence": True})
        self.assertEqual(template_choice_values(template, {"include_bond_sentence": "false"}), {"include_bond_sentence": False})


class ReviewedFieldQuestionsTests(SimpleTestCase):
    def template(self, schema, choices=()):
        block = SimpleNamespace(key="facts", label="Facts", body="It took {{ fields.discovery_days }} days and {{ fields.how_when }}.", ai_instructions=[], ai_fill_mode="revision_on_request")
        return SimpleNamespace(
            metadata={"fields": ["fields.discovery_days", "fields.how_when"], "fieldSchema": schema, "choices": list(choices)},
            blocks=SimpleNamespace(all=lambda: [block]),
        )

    def test_the_reviewers_question_and_kind_win(self):
        template = self.template({
            "discovery_days": {"label": "Days needed for discovery", "question": "How many days does discovery need?"},
            "how_when": {"label": "How the outage was discovered", "kind": "narrative", "question": "How and when was it discovered?"},
        })
        requests = {request.key: request for request in template_field_requests(template)}
        self.assertEqual(requests["discovery_days"].question, "How many days does discovery need?")
        self.assertEqual(requests["discovery_days"].label, "Days needed for discovery")
        self.assertEqual(requests["how_when"].kind, KIND_NARRATIVE)

    def test_a_field_written_up_under_an_alias_keeps_its_review(self):
        block = SimpleNamespace(key="facts", label="Facts", body="Rent is owed to {{ fields.landlord_name }}.", ai_instructions=[], ai_fill_mode="revision_on_request")
        template = SimpleNamespace(
            metadata={"fields": ["fields.landlord_name"], "fieldSchema": {"landlord_name": {"question": "Who is the landlord?", "optional": True}}, "choices": []},
            blocks=SimpleNamespace(all=lambda: [block]),
        )
        request = template_field_requests(template)[0]
        self.assertEqual(request.key, "plaintiff_name")
        self.assertEqual(request.question, "Who is the landlord?")
        self.assertTrue(request.optional)

    def test_a_field_for_a_choice_not_taken_is_not_asked(self):
        template = self.template(
            {"discovery_days": {"when": {"subsidy": "subsidized"}}},
            choices=[{"name": "subsidy", "options": ["subsidized", "not_subsidized"], "default": "subsidized"}],
        )
        self.assertTrue(field_applies(template, "discovery_days", {}))
        self.assertFalse(field_applies(template, "discovery_days", {"subsidy": "not_subsidized"}))


class MaintainedWordingReachesTheExportTests(SimpleTestCase):
    def test_a_section_nobody_changed_renders_from_the_word_original(self):
        source = "Served on {{ fields.service_date }}."
        rendered = "Served on May 1."
        self.assertEqual(_block_revision(rendered, source, rendered), "")
        # A reviewer's edit still reaches the export.
        self.assertEqual(_block_revision("Served on May 2 by hand.", source, rendered), "Served on May 2 by hand.")
        # Drafts composed before the template body was recorded behave as before.
        self.assertEqual(_block_revision(rendered, source), rendered)

    def test_a_maintained_statement_of_the_case_is_not_replaced_by_a_fact_list(self):
        block = SimpleNamespace(
            key="statement-of-the-case",
            label="Statement of the Case",
            block_type="facts",
            ai_fill_mode="revision_on_request",
            body="Plaintiff {{ fields.plaintiff_name }} is seeking to evict Defendant.",
            required=True,
            supporting_sources=[],
        )
        context = GenerationContext(
            matter=SimpleNamespace(summary="", jurisdiction="Housing Court", client_name="Jordan Client", external_id="1"),
            selected_facts=[],
            selected_curated_facts=[],
            selected_sources=[],
            template=SimpleNamespace(title="Motion", jurisdiction="", blocks=SimpleNamespace(all=lambda: [block])),
            mode="draft_from_template",
            template_data={"plaintiff_name": "Pat Landlord"},
        )

        section = ConstrainedDraftingService().compose_document(context, [block.key])[0]

        self.assertEqual(section["body"], "Plaintiff Pat Landlord is seeking to evict Defendant.")
        self.assertEqual(section["templateBody"], section["body"])


PRIVATE_TEMPLATES = organization_content_library_dir() / "document-templates"
REFINED = sorted(
    path.stem for path in (organization_content_library_dir() / "template-refinements").glob("*.yaml")
) if (organization_content_library_dir() / "template-refinements").exists() else []
# What a template author writes for the drafter rather than for the court.
EDITORIAL = re.compile(r"\[(?!s\])[^\]\n]{0,80}\]|\$X\b|\bX/X\b|\bmm/dd\b|\{\{|\{%|Case Caption|\(needs a cite\)")


class RefinedOrganizationTemplatesTests(TestCase):
    """Every refined organization template, answered in full, reads as a
    finished document. A checkout without the private provider skips this; it
    does not pass it."""

    def test_answered_templates_carry_no_editorial_markers(self):
        if not REFINED:
            self.skipTest("no organization template refinements in this checkout")
        checked = 0
        for slug in REFINED:
            manifest_path = PRIVATE_TEMPLATES / slug / "manifest.yaml"
            if not manifest_path.exists():
                continue
            manifest = yaml.safe_load(manifest_path.read_text())
            with self.subTest(template=slug):
                self.assertEqual((manifest["source"].get("refinements") or {}).get("path"), f"template-refinements/{slug}.yaml",
                                 "the package was not regenerated from its refinement file")
                answers = {path.removeprefix("fields."): f"Answer {index}" for index, path in enumerate(manifest["fields"])}
                for name, schema in (manifest.get("field_schema") or {}).items():
                    if schema.get("type") == "money":
                        answers[name] = "100"
                    if schema.get("type") == "list":
                        answers[name] = ["First Answer", "Second Answer"]
                context = {
                    "fields": template_field_values(answers),
                    "blocks": {block["key"]: {"revision": "", "paragraphs": ["Drafted."], "items": ["Drafted."]} for block in manifest["blocks"]},
                    "plaintiff": "Pat Landlord", "defendant": "Jordan Client", "client_name": "Jordan Client",
                    "client": {"name": "Jordan Client", "pronouns": ""}, "case_number": "2026-CVG-1",
                    "advocate_name": "Avery Advocate", "advocate_name_and_bar": "Avery Advocate (0012345)",
                    "advocate_signature": "/s/ Avery Advocate\nAvery Advocate (0012345)",
                    "advocate_signature_block": "/s/ Avery Advocate\nAvery Advocate (0012345)\nLegal Aid",
                    "advocate_phone": "216-555-0100", "advocate_email": "avery@example.org", "advocate_bar_number": "0012345",
                    **{choice["name"]: choice.get("default") for choice in manifest.get("choices", [])},
                    **{flag: True for flag in manifest.get("flags", [])},
                }
                with tempfile.TemporaryDirectory() as tmp:
                    output = Path(tmp) / "out.docx"
                    _render_docx_template(str(manifest_path.parent / "template.docx"), context, str(output))
                    document = Document(output)
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                text += "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                text += "\n".join(paragraph.text for section in document.sections for paragraph in section.header.paragraphs)
                self.assertEqual(EDITORIAL.findall(text), [])
                checked += 1
        self.assertTrue(checked, "refinement files exist but no package was prepared from them")


class DefaultAnswersTests(SimpleTestCase):
    def test_a_cleared_default_stays_cleared(self):
        from apps.drafting.services import _template_data_with_missing_information_answers

        cleared = _template_data_with_missing_information_answers({}, [
            {"field": "fields.judge", "answer": "", "has_default": True},
            {"field": "fields.magistrate", "answer": "", "not_needed": True},
        ])
        self.assertEqual(cleared, {"judge": ""})
        kept = _template_data_with_missing_information_answers({}, [{"field": "fields.judge", "answer": "W. Moná Scott", "has_default": True}])
        self.assertEqual(kept["judge"], "W. Moná Scott")
