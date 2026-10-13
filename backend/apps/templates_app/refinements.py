"""Hand-maintained conversion instructions for one maintained original.

Automatic conversion reads what the author marked as variable -- brackets,
underscores, highlighting -- and names each fill-in from the words around it.
That is the right default, and it cannot know what the author meant. The Word
originals carry editorial notes as well as blanks: "[written/oral]" is a choice,
"[X] days" is a number, "[etc.]" is a note to delete, and a highlighted sentence
may be wording to keep. Left to the heuristics, those print in a filing as
"[written/oral]" or bind the client's name under twenty different field names.

A refinement file records, for one template, what a reviewer decided those
passages mean. It lives beside the originals in the organization provider,
`template-refinements/<template-slug>.yaml`, and ingestion applies it to the
original *before* automatic conversion, so the converter still preserves every
run's formatting and still converts whatever the file does not mention.

Every edit names the exact wording it changes and how many times that wording
appears. An edit that no longer matches fails ingestion instead of being
skipped: an original revised in Word must not quietly lose the decision a
reviewer made about it, which would put "[OR]" back in a certificate of service
with nothing looking wrong.
"""

from __future__ import annotations

import hashlib
import re
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.templates_app.placeholders import Replacement, apply_replacements


REFINEMENTS_DIR = "template-refinements"
SCHEMA_VERSION = 1
FIELD_PATH_RE = re.compile(r"\bfields\.([A-Za-z_]\w*)")
FIELD_KINDS = {"value", "narrative"}
FIELD_TYPES = {"text", "date", "time", "money", "number", "list", "pronouns", "select", "paragraph"}
LATITUDES = {"locked", "guided", "generate"}


class RefinementError(ValueError):
    """A refinement file is malformed or no longer matches its original."""


@dataclass
class Refinements:
    path: Path
    data: dict
    sha256: str
    applied: list[str] = field(default_factory=list)

    @property
    def edits(self) -> list[dict]:
        return self.data.get("edits") or []

    @property
    def caption(self) -> dict | None:
        return self.data.get("caption") or None

    @property
    def choices(self) -> list[dict]:
        return self.data.get("choices") or []

    @property
    def field_schema(self) -> dict:
        return self.data.get("fields") or {}

    @property
    def block_overrides(self) -> dict:
        return self.data.get("blocks") or {}

    def manifest_overrides(self) -> dict:
        keys = ("title", "kind", "description", "goal", "negative_goal", "aliases", "jurisdiction")
        return {key: self.data[key] for key in keys if key in self.data}

    def referenced_fields(self) -> set[str]:
        """`fields.<name>` paths the file's own wording introduces."""
        texts = [edit.get("replace") or "" for edit in self.edits]
        caption = self.caption or {}
        texts += [str(line) for key in ("court", "left", "right") for line in caption.get(key) or []]
        return {f"fields.{name}" for text in texts for name in FIELD_PATH_RE.findall(text)}


def refinements_path(content_root: Path, slug: str) -> Path:
    return Path(content_root) / REFINEMENTS_DIR / f"{slug}.yaml"


def load_refinements(content_root: Path, slug: str) -> Refinements | None:
    path = refinements_path(content_root, slug)
    if not path.exists():
        return None
    raw = path.read_bytes()
    data = yaml.safe_load(raw) or {}
    digest = hashlib.sha256(raw)
    if data.get("extends"):
        # Settings several templates share -- one court's caption and its
        # judge -- live in one file, so a change of judge is one edit.
        base_path = path.parent / data["extends"]
        if not base_path.exists():
            raise RefinementError(f"{path}: extends {data['extends']!r}, which does not exist")
        base_raw = base_path.read_bytes()
        digest.update(base_raw)
        data = _merged(yaml.safe_load(base_raw) or {}, data)
    _validate(path, data, slug)
    return Refinements(path=path, data=data, sha256=digest.hexdigest())


def _merged(base, override):
    """The template's own settings over the shared file's, mapping by mapping;
    lists (edits, choices) run the shared entries first."""
    merged = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(base.get(key), dict):
            merged[key] = _merged(base[key], value)
        elif isinstance(value, list) and isinstance(base.get(key), list):
            merged[key] = [*base[key], *value]
        else:
            merged[key] = value
    return merged


def _validate(path: Path, data: dict, slug: str):
    if data.get("schema_version") != SCHEMA_VERSION:
        raise RefinementError(f"{path}: schema_version must be {SCHEMA_VERSION}")
    if data.get("slug") != slug:
        raise RefinementError(f"{path}: slug {data.get('slug')!r} does not name template {slug!r}")
    for index, edit in enumerate(data.get("edits") or [], start=1):
        if not isinstance(edit, dict) or not edit.get("find"):
            raise RefinementError(f"{path}: edit {index} needs `find`")
        unknown = set(edit) - {"find", "replace", "within", "count", "occurrence", "delete_paragraph", "note"}
        if unknown:
            raise RefinementError(f"{path}: edit {index} has unknown keys {sorted(unknown)}")
    for name, schema in (data.get("fields") or {}).items():
        if (schema or {}).get("kind", "value") not in FIELD_KINDS:
            raise RefinementError(f"{path}: field {name} kind must be one of {sorted(FIELD_KINDS)}")
        if "default" in (schema or {}) and not isinstance(schema["default"], str):
            raise RefinementError(f"{path}: field {name} default must be text")
        if (schema or {}).get("type", "text") not in FIELD_TYPES:
            raise RefinementError(f"{path}: field {name} type must be one of {sorted(FIELD_TYPES)}")
    for choice in data.get("choices") or []:
        values = [option["value"] if isinstance(option, dict) else option for option in choice.get("options") or []]
        if not choice.get("name") or len(values) < 2:
            raise RefinementError(f"{path}: a choice needs a name and at least two options")
        if choice.get("default", values[0]) not in values:
            raise RefinementError(f"{path}: choice {choice['name']} default is not one of its options")
    for key, override in (data.get("blocks") or {}).items():
        if (override or {}).get("ai_latitude", "locked") not in LATITUDES:
            raise RefinementError(f"{path}: block {key} ai_latitude must be one of {sorted(LATITUDES)}")


def _marker_paragraph(paragraph, text: str):
    marker = deepcopy(paragraph._p)
    for child in list(marker):
        if child.tag != qn("w:pPr"):
            marker.remove(child)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    marker.append(run)
    return marker


def _matches(text: str, find: str) -> list[int]:
    positions, start = [], 0
    while (position := text.find(find, start)) != -1:
        positions.append(position)
        start = position + len(find)
    return positions


def apply_edits(paragraphs, refinements: Refinements):
    """Rewrite the original's wording as the refinement file says.

    ``paragraphs`` is every paragraph of the document's story -- body, tables,
    headers and footers -- in document order. Edits apply in file order, each
    against the text the previous ones left. Substitution is run by run, so the
    wording around an edit keeps its formatting.

    Conditions are written inline (`{% if %}` inside the sentence) rather than
    as paragraph-level markers: the converter later wraps each block in its own
    paragraph-level revision conditional, and a marker paragraph inserted here
    would land on a block boundary and straddle it.
    """
    paragraphs = list(paragraphs)
    for index, edit in enumerate(refinements.edits, start=1):
        find = edit["find"]
        within = edit.get("within")
        hits = []
        for paragraph in paragraphs:
            if paragraph._p.getparent() is None:
                continue
            text = paragraph.text
            if within and within not in text:
                continue
            hits.extend((paragraph, position) for position in _matches(text, find))
        occurrence = edit.get("occurrence")
        expected = edit.get("count", 1)
        if occurrence is None and len(hits) != expected:
            raise RefinementError(
                f"{refinements.path}: edit {index} expected {expected} match(es) of {find!r}"
                f"{' within ' + repr(within) if within else ''}, found {len(hits)}"
            )
        if occurrence is not None:
            if not 1 <= occurrence <= len(hits):
                raise RefinementError(
                    f"{refinements.path}: edit {index} wants occurrence {occurrence} of {find!r}, found {len(hits)}"
                )
            hits = [hits[occurrence - 1]]

        by_paragraph: dict[int, tuple] = {}
        for paragraph, position in hits:
            by_paragraph.setdefault(id(paragraph._p), (paragraph, []))[1].append(position)
        for paragraph, positions in by_paragraph.values():
            if edit.get("delete_paragraph"):
                paragraph._p.getparent().remove(paragraph._p)
                continue
            if "replace" in edit:
                replacement = str(edit.get("replace") or "")
                apply_replacements(
                    paragraph,
                    [Replacement(position, position + len(find), replacement) for position in positions],
                )
        refinements.applied.append(find)


def _set_cell_borders(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        value = edges.get(edge)
        element.set(qn("w:val"), "single" if value else "nil")
        if value:
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)


def _fill_cell(cell, lines, *, bold_lines=(), center=False):
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.space_before = 0
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if line:
            run = paragraph.add_run(str(line))
            run.bold = (line in bold_lines) or None


def apply_caption(document, refinements: Refinements):
    """Replace a "Case Caption -- <title>" marker with the filing's caption.

    The Cleveland originals open with a highlighted note telling the advocate to
    paste the caption in, which a prepared template printed verbatim at the top
    of every filing. The caption is built as a table so it carries no body
    paragraphs of its own: block discovery, which reads body paragraphs, sees
    the document start where the author's wording does.
    """
    caption = refinements.caption
    if not caption:
        return
    marker = caption.get("marker", "")
    target = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip() and (paragraph.text.strip().startswith(marker) if marker else paragraph.text.strip().casefold().startswith("case caption"))
        ),
        None,
    )
    if target is None:
        raise RefinementError(f"{refinements.path}: no caption marker {marker or 'Case Caption'!r} in the original")

    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    court_cell = table.cell(0, 0).merge(table.cell(0, 1))
    _set_cell_borders(court_cell)
    court = [str(line) for line in caption.get("court") or []]
    _fill_cell(court_cell, [*court, ""], bold_lines=court, center=True)
    left, right = table.cell(1, 0), table.cell(1, 1)
    _set_cell_borders(left, right=True, bottom=True)
    _set_cell_borders(right)
    _fill_cell(left, caption.get("left") or [])
    title = str(caption.get("title") or "")
    right_lines = [*(caption.get("right") or []), "", title] if title else list(caption.get("right") or [])
    _fill_cell(right, right_lines, bold_lines=(title,))
    for paragraph in right.paragraphs:
        paragraph.paragraph_format.left_indent = 114300  # 0.125in, clear of the rule

    target._p.addprevious(table._tbl)
    spacer = _marker_paragraph(target, "")
    target._p.addprevious(spacer)
    target._p.getparent().remove(target._p)
