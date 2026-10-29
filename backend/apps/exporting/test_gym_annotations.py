import io
import tempfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document
from pypdf import PdfReader

from apps.argument_gym.models import GymChallenge, GymDocument, GymRun, GymWorkspace
from apps.argument_gym.testing import build_pdf
from apps.core.storage import get_document_storage


class GymAnnotationTests(TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        setting = override_settings(DOCUMENT_STORAGE_BACKEND="filesystem", DOCUMENT_STORAGE_ROOT=directory.name)
        setting.enable()
        self.addCleanup(setting.disable)
        self.user = User.objects.create_user(username="reviewer")
        self.client.force_login(self.user)
        self.workspace = GymWorkspace.objects.create(owner=self.user, title="Synthetic review")

    def upload(self, content, filename):
        response = self.client.post(f"/api/argument-gym/workspaces/{self.workspace.pk}/documents/", {
            "file": SimpleUploadedFile(filename, content),
        })
        self.assertEqual(response.status_code, 201, response.content)
        brief = GymDocument.objects.get(pk=response.json()["document"]["id"])
        self.run = GymRun.objects.create(workspace=self.workspace, brief=brief, status="complete")
        self.url = f"/api/argument-gym/runs/{self.run.pk}/artifacts/annotated/"
        return brief

    def challenge(self, quote, **kwargs):
        return GymChallenge.objects.create(run=self.run, target={"excerpt": quote},
            opponent_argument="Check the supporting record.", suggested_response="Add a verified record citation.", **kwargs)

    def test_docx_upload_download_preserves_formatting_existing_comments_and_tables(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("The notice ").bold = True
        p.add_run("was delivered on Monday.").italic = True
        doc.add_comment(p.runs, "Original reviewer note", author="Reviewer")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "A unique table passage."
        doc.add_paragraph("Repeated passage.")
        doc.add_paragraph("Repeated passage.")
        stream = io.BytesIO()
        doc.save(stream)
        original = stream.getvalue()
        brief = self.upload(original, "synthetic.docx")
        self.challenge("The notice was delivered on Monday.")
        self.challenge("A unique table passage.", disposition="addressed")
        self.challenge("Repeated passage.")
        self.challenge("Missing passage.", disposition="dismissed")
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        result = Document(io.BytesIO(response.content))
        self.assertEqual(len(result.comments), 5)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)
        notes = [c.text for c in result.comments]
        self.assertEqual(sum("UNPLACED" in c for c in notes), 2)
        self.assertTrue(any("addressed" in c for c in notes))
        with get_document_storage().open(brief.storage_key) as stored:
            self.assertEqual(stored.read(), original)
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_pdf_page_notes_preserve_pages_and_unplaced_are_explicit(self):
        self.upload(build_pdf(["First page text.", "Unique second page passage."]), "synthetic.pdf")
        self.challenge("Unique second page passage.")
        self.challenge("Not in the PDF.")
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        result = PdfReader(io.BytesIO(response.content))
        self.assertEqual(len(result.pages), 2)
        self.assertIn("First page text.", result.pages[0].extract_text())
        self.assertIn("UNPLACED", result.pages[0]["/Annots"][0].get_object()["/Contents"])
        note = result.pages[1]["/Annots"][0].get_object()
        self.assertEqual(note["/Subtype"], "/Text")
        self.assertIn("Page-level note", note["/Contents"])

    def test_access_status_missing_original_and_checksum(self):
        brief = self.upload(build_pdf(["Synthetic brief."]), "synthetic.pdf")
        self.challenge("Synthetic brief.")
        other = User.objects.create_user(username="other")
        self.client.force_login(other)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.client.force_login(self.user)
        self.assertEqual(self.client.post(self.url).status_code, 405)
        self.run.status = "running"
        self.run.save()
        self.assertEqual(self.client.get(self.url).status_code, 409)
        self.run.status = "complete"
        self.run.save()
        get_document_storage().put_bytes(content=b"changed", key=brief.storage_key)
        self.assertEqual(self.client.get(self.url).status_code, 409)
        brief.storage_key = ""
        brief.save()
        self.assertEqual(self.client.get(self.url).status_code, 409)

    def test_pdf_exhibits_are_preserved_but_never_used_as_comment_targets(self):
        self.upload(build_pdf(["Brief passage.", "Exhibit A", "Only in the exhibit."]), "filing.pdf")
        self.assertEqual(self.run.brief.page_range["end"], 1)
        self.challenge("Only in the exhibit.")
        response = self.client.get(self.url)
        result = PdfReader(io.BytesIO(response.content))
        self.assertEqual(len(result.pages), 3)
        self.assertIn("UNPLACED", result.pages[0]["/Annots"][0].get_object()["/Contents"])
        self.assertNotIn("/Annots", result.pages[2])

    def test_native_draft_export_and_stale_run_refusal(self):
        from apps.argument_gym.pipeline import brief_units, brief_snapshot
        from apps.drafting.models import DraftDocument, DraftingSession
        from apps.matters.models import Matter

        matter = Matter.objects.create(client_name="Synthetic client", external_id="gym-annotation-test", matter_type="Synthetic matter")
        session = DraftingSession.objects.create(matter=matter, mode="draft_from_scratch")
        draft = DraftDocument.objects.create(session=session, title="Synthetic brief", plain_text="A native draft passage.",
            sections=[{"key": "argument", "label": "Argument", "body": "A native draft passage."}])
        brief = GymDocument.objects.create(workspace=self.workspace, role="brief_under_test", source_type="draft_document", draft_document=draft, title=draft.title)
        brief_units(brief)
        self.run = GymRun.objects.create(workspace=self.workspace, brief=brief, status="complete", snapshot=brief_snapshot(brief))
        self.url = f"/api/argument-gym/runs/{self.run.pk}/artifacts/annotated/"
        self.challenge("A native draft passage.")
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200, response.content[:200])
        result = Document(io.BytesIO(response.content))
        self.assertEqual(len(result.comments), 1)
        self.assertNotIn("UNPLACED", list(result.comments)[0].text)
        draft.plain_text = "An edited passage."
        draft.save()
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 409)
        self.assertIn("changed since the run", response.json()["error"])
